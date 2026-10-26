"""Synthetic compositions: financial splits, OS unions, stock joins and provenance."""
import copy
import shutil
from dataclasses import asdict
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.core.batch import export_import
from app.core.catalog import verify_rules, workspace_config
from app.models import Cancelled, Control
from app.storage.imports import ImportStore
from app.storage.local import LocalStore
from app.storage.studio import StudioHistory
from app.storage.workspace import initialize, save_config, save_rule_revision
from app.studio.catalog import folder_tables, inventory
from app.studio.engine import execute
from app.studio.plan import approve, parse
from tests.conftest import ROOT


def library(tmp_path, make_xlsx, definitions):
    root, models = tmp_path / "library", tmp_path / "models"
    root.mkdir()
    models.mkdir()
    families = {}
    for name, headers, family in definitions:
        path = make_xlsx(headers, [], name)
        shutil.copy2(path, models / name)
        families[str(models / name)] = family
    save_config(root, {"models_dir": str(models), "families": families})
    save_rule_revision(root, "# Regras sintéticas\nPreservar fontes, códigos e contagens.", {"version": 1, "max_total_rows": 10000, "enabled_modules": [], "rules": []})
    return root


def project(tmp_path, paths, options, categories=None):
    store = ImportStore(LocalStore(tmp_path / "state"))
    categories = categories or ["generic"]
    inputs = [{"id": f"a{i}", "path": str(path), "title": path.stem, "family": categories[0], "role": "source", "header_row": 1}
              for i, path in enumerate(paths)]
    return store.create("Projeto sintético", categories, inputs, asdict(options), mode="studio")


def recipe(outputs, exclusions=None):
    return {"version": 1, "summary": "Plano sintético revisado", "questions": [], "outputs": outputs, "exclusions": exclusions or []}


def route(context, oid, target_name, source, mapping, *, filters=None, joins=None, constants=None):
    target = next(t for t in context["targets"] if Path(t["path"]).name == target_name)
    return {"id": oid, "title": oid, "target": target["id"], "group": "principal", "rules": [], "inputs": [
        {"source": context["sources"][source]["id"], "filters": filters or [], "joins": joins or [],
         "fields": [{"target": target, "source": origin} for target, origin in mapping.items()],
         "constants": [{"target": target, "value": value} for target, value in (constants or {}).items()]}]}


def ready(job, context, plan):
    job["studio"].update(plan=plan, approval=approve(plan, context))
    return job


def close(sessions):
    for session in sessions.values():
        session.close()


def test_packaged_rules_survive_deleted_onefile_directory(tmp_path):
    state = tmp_path / "state"
    root = initialize(ROOT, state)
    config = workspace_config(root)
    assert Path(config["profile_path"]).is_relative_to(root)
    vanished = tmp_path / "_MEI_old" / "app/core/rules/profile.json"
    config["profile_path"] = str(vanished)
    save_config(root, config)
    initialize(ROOT, state)
    assert verify_rules(root) and Path(workspace_config(root)["profile_path"]).is_file()
    # An arbitrary missing custom profile must not be replaced by bundled defaults.
    config = workspace_config(root)
    config["profile_path"] = str(tmp_path / "custom_missing.json")
    save_config(root, config)
    initialize(ROOT, state)
    with pytest.raises(ValueError, match="perfil executável"):
        verify_rules(root)


def test_combined_ctp_ctr_splits_and_keeps_zero_codes(tmp_path, make_xlsx, options):
    root = library(tmp_path, make_xlsx, [("ctp.xlsx", ["Documento", "Valor"], "pagar"), ("ctr.xlsx", ["Documento", "Valor"], "receber")])
    source = make_xlsx(["Documento", "Natureza", "Valor"], [["0001", "P", "10,00"], ["0002", "R", "20,00"]], "financeiro_junto.xlsx")
    job = project(tmp_path, [source], options, ["pagar", "receber"])
    context = inventory(root, job)
    outputs = [route(context, kind, kind + ".xlsx", 0, {"Documento": "C1", "Valor": "C3"}, filters=[{"column": "C2", "operator": "equals", "values": [value]}])
               for kind, value in [("ctp", "P"), ("ctr", "R")]]
    job, sessions = execute(root, tmp_path / "state", ready(job, context, recipe(outputs)))
    try:
        assert job["status"] == "Pronta para exportar"
        assert next(sessions["ctp"].rows())["values"]["Documento"] == "0001"
        assert next(sessions["ctr"].rows())["values"]["Documento"] == "0002"
        assert job["studio"]["last_run"]["unassigned_rows"] == 0
        result = export_import(job, sessions, tmp_path / "exports")
        assert len(list(result.rglob("*.xlsx"))) == 2
        assert len(list(result.rglob("linhagem.jsonl"))) == 2
        assert (result / "Parciais/plano_aprovado.json").is_file()
        workbook = load_workbook(next(result.rglob("pagar*.xlsx")))
        assert list(workbook.active.values)[1][0] == "0001"
        workbook.close()
        job["studio"]["instructions"] = "Uma instrução mudou"
        with pytest.raises(ValueError, match="mudaram"):
            export_import(job, sessions, tmp_path / "stale")
    finally:
        close(sessions)


def test_unknown_financial_row_and_overlapping_routes_block_export(tmp_path, make_xlsx, options):
    root = library(tmp_path, make_xlsx, [("ctp.xlsx", ["Documento"], "pagar"), ("ctr.xlsx", ["Documento"], "receber")])
    source = make_xlsx(["Documento", "Tipo"], [["001", "P"], ["002", "R"], ["003", "?"]])
    job = project(tmp_path, [source], options, ["pagar", "receber"])
    context = inventory(root, job)
    outputs = [route(context, kind, kind + ".xlsx", 0, {"Documento": "C1"}, filters=[{"column": "C2", "operator": "equals", "values": [value]}]) for kind, value in [("ctp", "P"), ("ctr", "R")]]
    result, sessions = execute(root, tmp_path / "state", ready(job, context, recipe(outputs)))
    try:
        assert result["studio"]["last_run"]["unassigned_rows"] == 1
        assert result["status"] == "Revisar pendências"
        with pytest.raises(ValueError):
            export_import(result, sessions, tmp_path / "blocked")
    finally:
        close(sessions)
    for out in outputs:
        out["inputs"][0]["filters"] = []
    result, sessions = execute(root, tmp_path / "state", ready(job, context, recipe(outputs)))
    try:
        assert any(i.code == "FINANCEIRO_DUPLICADO" for s in sessions.values() for i in s.global_issues)
    finally:
        close(sessions)


def test_product_stock_lookup_and_distinct_services_union(tmp_path, make_xlsx, options):
    root = library(tmp_path, make_xlsx, [("catalogo.xlsx", ["Código", "Descrição", "Quantidade"], "generic")])
    products = make_xlsx(["Código", "Nome"], [["0001", "Peça sintética"]], "produtos_recebidos.xlsx")
    stock = make_xlsx(["Produto", "Saldo"], [["0001", 5]], "estoque_recebido.xlsx")
    services = make_xlsx(["Código", "Nome"], [["S001", "Serviço sintético"]], "servicos_recebidos.xlsx")
    job = project(tmp_path, [products, stock, services], options)
    context = inventory(root, job)
    join = {"source": context["sources"][1]["id"], "alias": "estoque", "left": "C1", "right": "C1", "required": True}
    output = route(context, "catalogo", "catalogo.xlsx", 0, {"Código": "C1", "Descrição": "C2", "Quantidade": "estoque.C2"}, joins=[join])
    output["inputs"].extend(route(context, "x", "catalogo.xlsx", 2, {"Código": "C1", "Descrição": "C2"})["inputs"])
    result, sessions = execute(root, tmp_path / "state", ready(job, context, recipe([output])))
    try:
        rows = list(sessions["catalogo"].rows())
        assert rows[0]["values"]["Quantidade"] == "5" and rows[1]["values"]["Quantidade"] == ""
        assert len(list(sessions["catalogo"].lineage())[0]["joins"]) == 1
        assert result["studio"]["last_run"]["unassigned_rows"] == 0
        stock.write_bytes(stock.read_bytes() + b"changed")
        with pytest.raises(ValueError, match="fontes compostas"):
            sessions["catalogo"].verify()
    finally:
        close(sessions)


def test_duplicate_lookup_key_fails_without_multiplication(tmp_path, make_xlsx, options):
    root = library(tmp_path, make_xlsx, [("destino.xlsx", ["Código", "Saldo"], "generic")])
    left = make_xlsx(["Código"], [["001"]], "produtos.xlsx")
    right = make_xlsx(["Código", "Saldo"], [["001", 1], ["001", 2]], "estoque.xlsx")
    job = project(tmp_path, [left, right], options)
    context = inventory(root, job)
    join = {"source": context["sources"][1]["id"], "alias": "stock", "left": "C1", "right": "C1", "required": True}
    output = route(context, "produto", "destino.xlsx", 0, {"Código": "C1", "Saldo": "stock.C2"}, joins=[join])
    with pytest.raises(ValueError, match="duplicada"):
        execute(root, tmp_path / "state", ready(job, context, recipe([output])))
    assert not list((tmp_path / "state/sessions").glob("*.sqlite3"))


def test_os_headers_parts_and_services_become_linked_outputs(tmp_path, make_xlsx, options):
    root = library(tmp_path, make_xlsx, [("ordens_somente_cabecalhos.xlsx", ["Número da OS (externo)"], "os"),
                                      ("ordens_somente_itens.xlsx", ["Número da OS (externo)", "Produto/Serviço (identificação)"], "os")])
    header = make_xlsx(["OS"], [["0001"]], "os_origem.xlsx")
    parts = make_xlsx(["OS", "Peça"], [["0001", "P001"]], "pecas_origem.xlsx")
    services = make_xlsx(["OS", "Serviço"], [["0001", "S001"]], "servicos_origem.xlsx")
    job = project(tmp_path, [header, parts, services], options, ["os"])
    context = inventory(root, job)
    main = route(context, "os", "ordens_somente_cabecalhos.xlsx", 0, {"Número da OS (externo)": "C1"})
    items = route(context, "itens", "ordens_somente_itens.xlsx", 1, {"Número da OS (externo)": "C1", "Produto/Serviço (identificação)": "C2"})
    items["inputs"].extend(route(context, "x", "ordens_somente_itens.xlsx", 2, {"Número da OS (externo)": "C1", "Produto/Serviço (identificação)": "C2"})["inputs"])
    result, sessions = execute(root, tmp_path / "state", ready(job, context, recipe([main, items])))
    try:
        assert result["status"] == "Pronta para exportar"
        assert sessions["itens"].stats()["read"] == 2
        export_import(result, sessions, tmp_path / "out")
    finally:
        close(sessions)


def test_recipe_is_data_not_code_and_approval_is_snapshot_bound(tmp_path, make_xlsx, options):
    root = library(tmp_path, make_xlsx, [("destino.xlsx", ["Nome"], "generic")])
    job = project(tmp_path, [make_xlsx(["Nome"], [["Sintético"]])], options)
    context = inventory(root, job)
    plan = recipe([route(context, "out", "destino.xlsx", 0, {"Nome": "C1"})])
    malicious = {**plan, "python": "delete everything"}
    with pytest.raises(ValueError, match="campos"):
        parse(malicious, context)
    bad = copy.deepcopy(plan)
    bad["outputs"][0]["inputs"][0]["fields"][0]["source"] = "eval(payload)"
    with pytest.raises(ValueError, match="inexistente"):
        parse(bad, context)
    ready(job, context, plan)
    job["studio"]["instructions"] = "changed"
    with pytest.raises(ValueError, match="mudaram"):
        execute(root, tmp_path / "state", job)
    with pytest.raises(Cancelled):
        execute(root, tmp_path / "state", job, Control(cancelled=lambda: True))


def test_folder_selection_and_append_only_history(tmp_path, make_xlsx):
    make_xlsx(["Nome"], [["Sintético"]], "fonte.xlsx")
    (tmp_path / "~$ignorar.xlsx").write_text("not a source")
    found = folder_tables(tmp_path)
    assert len(found) == 1 and found[0]["source_sheet"] == "Dados"
    store = LocalStore(tmp_path / "state")
    history = StudioHistory(store)
    first = history.append("job", "user", {"text": "Separe contas"})
    second = history.append("job", "assistant", {"text": "Qual coluna?"})
    assert first < second and len(StudioHistory(store).list("job")) == 2


def test_empty_destination_and_justified_exclusion_are_explicit(tmp_path, make_xlsx, options):
    root = library(tmp_path, make_xlsx, [("ctp.xlsx", ["Documento"], "pagar"), ("ctr.xlsx", ["Documento"], "receber")])
    job = project(tmp_path, [make_xlsx(["Documento", "Tipo"], [["001", "P"], ["002", "IGNORAR"]])], options, ["pagar", "receber"])
    context = inventory(root, job)
    outputs = [route(context, name, name + ".xlsx", 0, {"Documento": "C1"}, filters=[{"column": "C2", "operator": "equals", "values": [kind]}]) for name, kind in [("ctp", "P"), ("ctr", "R")]]
    exclusion = {"source": context["sources"][0]["id"], "filters": [{"column": "C2", "operator": "equals", "values": ["IGNORAR"]}], "reason": "Exclusão sintética expressamente autorizada"}
    result, sessions = execute(root, tmp_path / "state", ready(job, context, recipe(outputs, [exclusion])))
    try:
        assert result["status"] == "Pronta para exportar"
        assert sessions["ctr"].stats()["read"] == 0
        assert result["studio"]["last_run"]["exclusions"][0]["rows"] == 1
        path = export_import(result, sessions, tmp_path / "out")
        book = load_workbook(next(path.rglob("receber*.xlsx")))
        assert list(book.active.values) == [("Documento",)]
        book.close()
    finally:
        close(sessions)


def test_constants_are_audited_and_missing_required_lookup_blocks(tmp_path, make_xlsx, options):
    root = library(tmp_path, make_xlsx, [("destino.xlsx", ["Código", "Classe", "Saldo"], "generic")])
    products = make_xlsx(["Código"], [["001"], ["002"]], "produtos.xlsx")
    stock = make_xlsx(["Código", "Saldo"], [["001", 5]], "estoque.xlsx")
    job = project(tmp_path, [products, stock], options)
    context = inventory(root, job)
    join = {"source": context["sources"][1]["id"], "alias": "stock", "left": "C1", "right": "C1", "required": True}
    output = route(context, "out", "destino.xlsx", 0, {"Código": "C1", "Saldo": "stock.C2"}, constants={"Classe": "Serviço"}, joins=[join])
    result, sessions = execute(root, tmp_path / "state", ready(job, context, recipe([output])))
    try:
        assert sessions["out"].stats()["errors"] == 1
        assert next(sessions["out"].rows())["changes"][0]["rule"] == "PLANO_CONSTANTE_APROVADA"
        assert next(sessions["out"].lineage())["operations"]["constant_fields"] == ["Classe"]
        with pytest.raises(ValueError):
            export_import(result, sessions, tmp_path / "blocked")
    finally:
        close(sessions)


def test_unresolved_questions_and_duplicate_json_keys_cannot_be_approved(tmp_path, make_xlsx, options):
    root = library(tmp_path, make_xlsx, [("destino.xlsx", ["Nome"], "generic")])
    job = project(tmp_path, [make_xlsx(["Nome"], [["Sintético"]])], options)
    context = inventory(root, job)
    plan = recipe([route(context, "out", "destino.xlsx", 0, {"Nome": "C1"})])
    plan["questions"] = ["Qual é a regra de classificação?"]
    with pytest.raises(ValueError, match="perguntas"):
        approve(plan, context)
    with pytest.raises(ValueError, match="duplicadas"):
        parse('{"version":1,"version":2}', context)


def test_correction_cannot_move_to_another_record_after_filter_change(tmp_path, make_xlsx, options):
    from app.studio.engine import row_identity
    root = library(tmp_path, make_xlsx, [("destino.xlsx", ["Nome"], "generic")])
    source = make_xlsx(["Nome", "Tipo"], [["Pessoa A", "P"], ["Pessoa B", "R"]])
    job = project(tmp_path, [source], options)
    context = inventory(root, job)
    output = route(context, "out", "destino.xlsx", 0, {"Nome": "C1"})
    plan = recipe([output])
    _, sessions = execute(root, tmp_path / "state", ready(job, context, plan))
    row = next(sessions["out"].rows())
    origin = next(sessions["out"].lineage())
    identity = row_identity(origin["source_id"], origin["source_line"], origin["joins"], row["raw"])
    close(sessions)
    job["studio"]["corrections"] = {"out": [{"line": 2, "field": "Nome", "value": "Pessoa A revisada", "origin_hash": identity}]}
    # Keep only the second physical row; its new output position is still 2.
    output["inputs"][0]["filters"] = [{"column": "C2", "operator": "equals", "values": ["R"]}]
    exclusion = {"source": context["sources"][0]["id"], "filters": [{"column": "C2", "operator": "equals", "values": ["P"]}], "reason": "Exclusão sintética da linha A"}
    context = inventory(root, job)
    _, sessions = execute(root, tmp_path / "state", ready(job, context, recipe([output], [exclusion])))
    try:
        row = next(sessions["out"].rows())
        assert row["values"]["Nome"] == "Pessoa B"
        assert row["severity"] == "ERRO" and "desatualizada" in row["issues"][0]["message"]
    finally:
        close(sessions)


def test_project_field_rules_are_executable_and_counted_in_the_audit(tmp_path, make_xlsx, options):
    root = library(tmp_path, make_xlsx, [("destino.xlsx", ["Nome", "Data"], "generic")])
    job = project(tmp_path, [make_xlsx(["Nome", "Data"], [["pessoa sintética", "2026-10-26"]])], options)
    context = inventory(root, job)
    output = route(context, "out", "destino.xlsx", 0, {"Nome": "C1", "Data": "C2"})
    def rule(ident, field, operation, input_format=""):
        return {"id": ident, "field": field, "operation": operation, "input_format": input_format,
                "value": "", "values": [], "min": 0, "max": 32767, "severity": "ERRO"}
    output["rules"] = [rule("NOME_MAIUSCULO", "Nome", "uppercase"), rule("DATA_CONFIRMADA", "Data", "date", "%Y-%m-%d")]
    result, sessions = execute(root, tmp_path / "state", ready(job, context, recipe([output])))
    try:
        assert next(sessions["out"].rows())["values"] == {"Nome": "PESSOA SINTÉTICA", "Data": "26/10/2026"}
        assert result["studio"]["last_run"]["outputs"]["out"]["transformations"] == {"NOME_MAIUSCULO": 1, "DATA_CONFIRMADA": 1}
    finally:
        close(sessions)
