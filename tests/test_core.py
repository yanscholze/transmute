import json

import pytest
from openpyxl import load_workbook

from app.core.catalog import digest, verify_rules
from app.core.exporters.writer import export, plan_parts
from app.core.importers.reader import excel_cell, inspect, records, sheets
from app.core.mapping import suggest, validate_mapping
from app.core.session import process, validate_links
from app.core.validation.values import decimal_value, document_valid
from app.models import Cancelled, Control
from tests.conftest import ROOT, synthetic_cpf


def test_originals_unchanged():
    expected = json.loads((ROOT / "docs/originais.sha256.json").read_text())
    assert all(digest(ROOT / path) == value for path, value in expected.items())


def test_layout_versions_and_headers(layouts):
    assert len(layouts) == 36
    assert len({x.digest for x in layouts}) == 12
    assert next(x for x in layouts if x.family == "clientes").headers[2] == "CPF/CNPJ"


def test_rules_drift_blocks(tmp_path):
    (tmp_path / "INSTRUCOES_CONVERSOES.md").write_text("regras alteradas")
    with pytest.raises(ValueError, match="mudou"):
        verify_rules(tmp_path)


def test_text_zeros_and_cpf():
    cpf = synthetic_cpf()
    assert cpf.startswith("0") and document_valid(cpf)
    assert not document_valid("11111111111")
    assert not document_valid("123")
    assert excel_cell("000001").text == "000001"
    assert excel_cell(123, "n", "000000").text == "000123"
    assert excel_cell(1e18, "n").text == "1000000000000000000"


@pytest.mark.parametrize("encoding,delimiter", [("utf-8-sig", ";"), ("utf-16", "\t"), ("cp1252", ";")])
def test_csv_encodings(tmp_path, encoding, delimiter):
    path = tmp_path / "entrada.csv"
    path.write_text(f'Nome{delimiter}Código\n"João ç"{delimiter}000123\n', encoding=encoding)
    source = inspect(path, header_row=1)
    row = next(records(source))[1]
    assert source.delimiter == delimiter
    assert row["C2"].text == "000123"
    assert row["C1"].text == "João ç"


def test_csv_ragged_and_blank(tmp_path):
    path = tmp_path / "entrada.csv"
    path.write_text("A;B\n\n1;2;3\n4\n", encoding="utf-8")
    source = inspect(path, header_row=1, delimiter=";")
    data = list(records(source))
    assert data[0][1] is None
    assert data[1][2] and data[2][2]
    assert data[1][1]["C3"].text == "3"


def test_html_disguised_xls(tmp_path):
    path = tmp_path / "entrada.xls"
    path.write_text('<html><table><tr><td>Código</td><td>Descrição</td></tr><tr><td>000001</td><td>Peça ç</td></tr></table></html>', encoding="utf-8")
    source = inspect(path, header_row=1)
    assert source.format == "html"
    assert next(records(source))[1]["C1"].text == "000001"


def test_multiple_sheets_and_header(make_xlsx):
    path = make_xlsx(["Relatório sintético"], [["Nome", "Código"], ["Item A", "0001"]])
    book = load_workbook(path)
    book.create_sheet("Outra")
    book.save(path)
    book.close()
    assert sheets(path) == ["Dados", "Outra"]
    source = inspect(path)
    assert source.header_row == 2


def test_duplicate_headers_not_auto_confirmed(make_xlsx, layout):
    source = inspect(make_xlsx(["CPF/CNPJ", "CPF/CNPJ"], [["1", "2"]]), header_row=1)
    mapping = suggest(source, layout("clientes"))
    assert all(not d.confirmed for d in mapping)
    assert source.warnings
    with pytest.raises(ValueError):
        validate_mapping(source, layout("clientes"), mapping)


def test_product_case_sensitive_mapping(make_xlsx, layout):
    source = inspect(make_xlsx(["Tipo do item", "Tipo do Item"], [["Produto", "00"]]), header_row=1)
    mapping = suggest(source, layout("produtos"))
    assert mapping[0].target == "Tipo do item"
    assert mapping[1].target == "Tipo do Item"


def test_clients_preservation_defaults_and_invalid_contacts(convert, layout):
    cpf = synthetic_cpf()
    session = convert(layout("clientes"), ["Nome", "CPF", "UF", "E-mail", "Telefone"], [["Pessoa Sintética ç", cpf, "SC", "x", "."]])
    row = next(session.rows())
    assert row["values"]["CPF/CNPJ"] == cpf
    assert row["values"]["Regime Tributário"] == "Normal"
    assert row["values"]["Ativo"] == ""
    assert row["values"]["E-mail"] == "" and row["severity"] == "AVISO"
    assert session.stats()["read"] == 1


def test_cpf_ie_not_guessed(convert, layout):
    session = convert(layout("clientes"), ["CPF", "Inscrição Estadual", "UF"], [[synthetic_cpf(), "000123", "SC"]])
    assert session.stats()["errors"] == 1
    row = next(session.rows())
    assert row["values"]["Inscrição Estadual"] == "000123"


def test_rg_composite_and_manual_correction(convert, layout, options):
    session = convert(layout("clientes"), ["CPF", "IE/RG", "UF"], [[synthetic_cpf(), "000123", "SC"]])
    assert next(session.rows())["values"]["RG"] == "000123"
    options.overrides = {(2, "UF"): "PR"}
    other = convert(layout("clientes"), ["CPF", "UF"], [[synthetic_cpf(), ""]], options, "outro.xlsx")
    assert next(other.rows())["values"]["UF"] == "PR"
    assert any(c["rule"] == "REVISAO_USUARIO" for c in next(other.rows())["changes"])


def test_missing_uf_and_confirmation(convert, layout, options):
    session = convert(layout("clientes"), ["Nome"], [["Pessoa Sintética"]])
    assert session.stats()["errors"] == 1
    options.geography_confirmed = False
    session = convert(layout("fornecedores"), ["UF"], [["SC"]], options, "fornecedor.xlsx")
    assert session.stats()["errors"] == 1


def test_active_requires_per_conversion_answer(make_xlsx, layout, options, tmp_path):
    path = make_xlsx(["UF"], [["SC"]])
    source = inspect(path, header_row=1)
    options.active = ""
    with pytest.raises(ValueError, match="Ativo"):
        process(ROOT, tmp_path, source, layout("clientes"), suggest(source, layout("clientes")), options)


def test_service_tax_and_not_generic_product_code(convert, layout):
    service = convert(layout("produtos"), ["Descrição", "Tipo do item"], [["Serviço sintético", "Serviço"]])
    v = next(service.rows())["values"]
    assert v["Tipo do Item"] == "09" and v["CFOP Dentro"] == "5933" and v["Situação tributária"] == "SERVIÇOS"
    product = convert(layout("produtos"), ["Descrição", "Tipo do item"], [["Peça sintética", "Produto"]], name="produto.xlsx")
    assert product.stats()["errors"] == 1
    assert next(product.rows())["values"]["Tipo do Item"] == ""


def test_oil_parts_not_classified(convert, layout):
    product = convert(layout("produtos"), ["Descrição", "Tipo do item", "Tipo do Item"], [["Filtro de óleo sintético", "Produto", "00"]])
    v = next(product.rows())["values"]
    assert v["Situação tributária"] == "VENDA NORMAL"
    assert v["CFOP Dentro"] == "5.102"


def test_km_letters_not_stripped(convert, layout):
    session = convert(layout("os", "joined"), ["Número da OS (externo)", "Quilometragem"], [["000001", "MODELO TR4"]])
    assert next(session.rows())["values"]["Quilometragem"] == "MODELO TR4"
    assert session.stats()["errors"] == 1


def test_numeric_identifier_and_formula_block(convert, layout, tmp_path):
    session = convert(layout("veiculos"), ["id_veiculo", "obs"], [[123, "=1+1"]])
    assert session.stats()["errors"] == 1
    with pytest.raises(ValueError, match="bloqueada"):
        export([session], tmp_path)
    report = export([session], tmp_path, reports_only=True)
    assert (report / "relatorio.json").is_file()
    assert not list(report.glob("*.xlsx"))


def test_export_zeros_order_omit_and_no_original_change(convert, layout, tmp_path):
    session = convert(layout("clientes"), ["Nome", "CPF", "UF"], [["Pessoa Sintética", synthetic_cpf(), "SC"]])
    before = digest(session.source.path)
    folder = export([session], tmp_path)
    path = next(folder.glob("*.xlsx"))
    book = load_workbook(path)
    sheet = book.active
    assert list(sheet.values)[0] == ("Razão Social / Nome", "CPF/CNPJ", "Regime Tributário", "UF")
    assert sheet.cell(2, 2).data_type == "s" and sheet.cell(2, 2).value == synthetic_cpf()
    assert digest(session.source.path) == before
    assert (folder / "diagnostico.md").exists() and (folder / "resumo.txt").exists()
    assert export([session], tmp_path) != folder
    book.close()


def test_large_csv_split_and_counts(tmp_path, layout, options):
    path = tmp_path / "grande.csv"
    with path.open("w") as stream:
        stream.write("descricao;id_veiculo\n")
        for i in range(20005):
            stream.write(f"Veículo sintético {i};{i:08d}\n")
    source = inspect(path, header_row=1)
    session = process(ROOT, tmp_path, source, layout("veiculos"), suggest(source, layout("veiculos")), options)
    try:
        folder = export([session], tmp_path)
        books = [load_workbook(p, read_only=True) for p in sorted(folder.glob("*.xlsx"))]
        counts = [sum(1 for _ in b.active.rows) - 1 for b in books]
        assert counts == [9999, 9999, 7]
        assert session.stats()["read"] == sum(counts)
        for b in books:
            b.close()
    finally:
        session.close()


def test_links_and_joint_partition(convert, layout, tmp_path):
    headers = convert(layout("os", "header"), ["Número da OS (externo)"], [["000001"], ["000002"]], name="cab.xlsx")
    items = convert(layout("os", "items"), ["Número da OS (externo)", "Quantidade"], [["000001", "1"], ["000001", "2"], ["000002", "1"]], name="itens.xlsx")
    assert not validate_links(headers, items)
    assignment, parts = plan_parts([headers, items], [None, None], 2)
    assert parts == 2 and assignment == {"000001": 1, "000002": 2}
    folder = export([headers, items], tmp_path)
    assert len(list(folder.glob("*.xlsx"))) == 2
    assert validate_links(items)[0].severity == "ERRO"


def test_oversized_group_and_interleaving_block(convert, layout):
    items = convert(layout("os", "joined"), ["Número da OS (externo)"], [["1"], ["1"], ["1"]])
    with pytest.raises(ValueError, match="único grupo"):
        plan_parts([items], [None], 2)
    other = convert(layout("os", "joined"), ["Número da OS (externo)"], [["1"], ["2"], ["1"]], name="intercalado.xlsx")
    with pytest.raises(ValueError, match="intercalados"):
        plan_parts([other], [None], 2)


def test_cancel_processing_and_no_session_leak(tmp_path, make_xlsx, layout, options):
    source = inspect(make_xlsx(["descricao"], [["Sintético"]]), header_row=1)
    with pytest.raises(Cancelled):
        process(ROOT, tmp_path, source, layout("veiculos"), suggest(source, layout("veiculos")), options, Control(cancelled=lambda: True))
    assert not list((tmp_path / "sessions").glob("*.sqlite3"))


def test_source_mutation_blocks_export(convert, layout, tmp_path):
    session = convert(layout("veiculos"), ["descricao"], [["Sintético"]])
    session.source.path.write_bytes(b"origem alterada")
    with pytest.raises(ValueError, match="origem mudou"):
        export([session], tmp_path)


def test_decimal_no_rounding_or_ambiguous_guess():
    assert str(decimal_value("1.234,56789")) == "1234.56789"
    with pytest.raises(ValueError, match="ambíguo"):
        decimal_value("1.234")
