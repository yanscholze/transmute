"""Trabalhos, processamento em lote e PDF usam somente arquivos sintéticos."""
from dataclasses import asdict
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.core.batch import export_import, process_import
from app.core.importers.reader import inspect, records, sheets
from app.models import Cancelled, Control
from app.storage.imports import ImportStore
from app.storage.local import LocalStore
from tests.conftest import ROOT

FIXTURES = Path(__file__).parent / "fixtures"


def attachment(layout, path, identifier="source"):
    return {"id": identifier, "path": str(path), "title": path.stem, "family": layout.family,
            "layout_path": str(layout.path), "layout_sheet": layout.sheet, "role": "source", "group": "principal", "status": "Aguardando"}


def test_pdf_tables_and_zero_preservation():
    path = FIXTURES / "tabelas_sinteticas.pdf"
    assert sheets(path) == ["Página 1 · Tabela 1", "Página 2 · Tabela 1"]
    source = inspect(path, sheets(path)[0], header_row=1)
    assert source.format == "pdf" and source.warnings
    assert next(records(source))[1]["C1"].text == "000001"
    assert "Página 1" in next(records(source))[1]["C1"].note


def test_pdf_without_table_is_explicitly_rejected():
    with pytest.raises(ValueError, match="sem tabela"):
        sheets(FIXTURES / "sem_tabela_sintetico.pdf")


def test_named_job_multiple_categories_export(tmp_path, make_xlsx, layout, options):
    local = LocalStore(tmp_path / "state")
    store = ImportStore(local)
    car = attachment(layout("veiculos"), make_xlsx(["descricao"], [["Veículo sintético"]], "veiculos.xlsx"), "cars")
    service = attachment(layout("produtos"), make_xlsx(["Descrição", "Tipo do item"], [["Serviço sintético", "Serviço"]], "servicos.xlsx"), "services")
    job = store.create("Importação sintética", ["veiculos", "produtos"], [car, service], asdict(options))
    job, sessions = process_import(ROOT, tmp_path, job)
    try:
        assert job["status"] == "Pronta para exportar" and len(sessions) == 2
        store.save(job)
        assert store.get(job["id"])["name"] == "Importação sintética"
        output = export_import(job, sessions, tmp_path / "saida")
        assert len(list(output.rglob("*.xlsx"))) == 2
        assert len(list((output / "Tratadas").rglob("*.xlsx"))) == 2
        assert len(list((output / "Parciais").rglob("relatorio.json"))) == 2
        assert (output / "importacao.json").is_file()
        assert not list((tmp_path / "saida").glob(".importacao-*"))
    finally:
        for session in sessions.values():
            session.close()


def test_every_pdf_table_represented_and_requires_review(tmp_path, layout, options):
    store = ImportStore(LocalStore(tmp_path))
    item = attachment(layout("veiculos"), FIXTURES / "tabelas_sinteticas.pdf")
    job = store.create("PDF sintético", ["veiculos"], [item], asdict(options))
    job, sessions = process_import(ROOT, tmp_path, job)
    try:
        assert len(job["attachments"]) == 2 and len(sessions) == 2
        assert job["status"] == "Revisar pendências"
        with pytest.raises(ValueError, match="pendentes"):
            export_import(job, sessions, tmp_path / "bloqueado")
        report = export_import(job, sessions, tmp_path / "relatorios", reports_only=True)
        assert (report / "importacao.json").is_file()
        assert not list(report.rglob("*.xlsx"))
        diagnostics = list((report / "Parciais").rglob("diagnostico.md"))
        assert len(diagnostics) == 2
        assert all("PDF_CONFERENCIA" in path.read_text() for path in diagnostics)
    finally:
        for session in sessions.values():
            session.close()
    job["options"]["pdf_confirmed"] = True
    job, sessions = process_import(ROOT, tmp_path, job)
    try:
        assert job["status"] == "Pronta para exportar"
        out = export_import(job, sessions, tmp_path / "conferido")
        assert len(list(out.rglob("*.xlsx"))) == 2
    finally:
        for session in sessions.values():
            session.close()


def test_failed_file_stays_in_import(tmp_path, layout, make_xlsx, options):
    good = attachment(layout("veiculos"), make_xlsx(["descricao"], [["Sintético"]]), "good")
    pdf = attachment(layout("veiculos"), FIXTURES / "sem_tabela_sintetico.pdf", "pdf")
    store = ImportStore(LocalStore(tmp_path))
    job = store.create("Pendências", ["veiculos"], [good, pdf], asdict(options))
    job, sessions = process_import(ROOT, tmp_path, job)
    try:
        assert len(job["attachments"]) == 2 and len(sessions) == 1
        assert job["attachments"][1]["status"] == "Precisa de revisão"
        assert "sem tabela" in job["attachments"][1]["message"]
    finally:
        for session in sessions.values():
            session.close()


def test_multi_sheet_and_changed_header_does_not_reuse_mapping(tmp_path, layout, make_xlsx, options):
    path = make_xlsx(["descricao"], [["Sintético"]])
    book = load_workbook(path)
    other = book.create_sheet("Outra")
    other.append(["descricao"])
    other.append(["Segundo veículo sintético"])
    book.save(path)
    book.close()
    store = ImportStore(LocalStore(tmp_path))
    job = store.create("Duas abas", ["veiculos"], [attachment(layout("veiculos"), path)], asdict(options))
    job, sessions = process_import(ROOT, tmp_path, job)
    for session in sessions.values():
        session.close()
    assert len(job["attachments"]) == 2
    book = load_workbook(path)
    book["Dados"]["A1"] = "campo alterado"
    book.save(path)
    book.close()
    job, sessions = process_import(ROOT, tmp_path, job)
    try:
        assert "estrutura da origem mudou" in job["attachments"][0]["message"]
        assert "stats" not in job["attachments"][0]
    finally:
        for session in sessions.values():
            session.close()


def test_saved_corrections_and_cancel_do_not_publish(tmp_path, layout, make_xlsx, options):
    item = attachment(layout("veiculos"), make_xlsx(["descricao"], [["Antes sintético"]]))
    item["corrections"] = [{"line": 2, "field": "descricao", "value": "Depois sintético"}]
    store = ImportStore(LocalStore(tmp_path))
    job = store.create("Correções", ["veiculos"], [item], asdict(options))
    job, sessions = process_import(ROOT, tmp_path, job)
    try:
        assert next(sessions["source"].rows())["values"]["descricao"] == "Depois sintético"
        with pytest.raises(Cancelled):
            export_import(job, sessions, tmp_path / "saida", Control(cancelled=lambda: True))
        assert not list((tmp_path / "saida").iterdir())
    finally:
        for session in sessions.values():
            session.close()


def test_batch_links_and_ambiguous_pairs_remain_pending(tmp_path, layout, make_xlsx, options):
    """A lista de arquivos não autoriza adivinhar qual cabeçalho recebe itens."""
    header = attachment(layout("os", "header"), make_xlsx(["Número da OS (externo)"], [["000001"]], "cab.xlsx"), "header")
    items = attachment(layout("os", "items"), make_xlsx(["Número da OS (externo)", "Quantidade"], [["000001", "2"]], "itens.xlsx"), "items")
    store = ImportStore(LocalStore(tmp_path))
    job = store.create("OS sintética", ["os"], [header, items], asdict(options))
    job, sessions = process_import(ROOT, tmp_path, job)
    try:
        assert job["status"] == "Pronta para exportar"
        output = export_import(job, sessions, tmp_path / "saida")
        assert len(list((output / "Tratadas").rglob("*.xlsx"))) == 2
    finally:
        for session in sessions.values():
            session.close()
    extra = {**header, "id": "duplicate", "title": "Segundo cabeçalho"}
    job["attachments"].append(extra)
    job, sessions = process_import(ROOT, tmp_path, job)
    try:
        assert job["status"] == "Revisar pendências" and "par único" in job["link_message"]
        with pytest.raises(ValueError, match="par único"):
            export_import(job, sessions, tmp_path / "bloqueado")
        report = export_import(job, sessions, tmp_path / "relatorio", reports_only=True)
        assert len(list(report.rglob("relatorio.json"))) == 3
        assert not list((tmp_path / "bloqueado").glob(".importacao-*"))
    finally:
        for session in sessions.values():
            session.close()
