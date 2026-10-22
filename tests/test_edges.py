import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.ai.cli import CLIProvider
from app.core.exporters.writer import export
from app.core.importers.reader import inspect, records
from app.core.mapping import suggest
from app.core.session import process, validate_links
from app.models import Cancelled, Control
from tests.conftest import ROOT


def test_real_biff_xls_format_zeros():
    path = Path(__file__).parent / "fixtures/sintetico_biff.xls"
    source = inspect(path, header_row=1)
    assert source.format == "xls" and source.sheet == "Dados"
    row = next(records(source))[1]
    assert row["C1"].text == "000123" and row["C1"].kind == "formatted_number"
    assert row["C2"].text == "Peça sintética ç"


def test_formula_like_csv_exported_as_literal(tmp_path, layout, options):
    path = tmp_path / "formula.csv"
    path.write_text('descricao;obs\nSintético;=HYPERLINK("https://example.invalid")\n', encoding="utf-8")
    source = inspect(path, header_row=1)
    session = process(ROOT, tmp_path, source, layout("veiculos"), suggest(source, layout("veiculos")), options)
    try:
        folder = export([session], tmp_path)
        book = load_workbook(next(folder.glob("*.xlsx")))
        try:
            assert book.active.cell(2, 2).data_type == "s"
            assert book.active.cell(2, 2).value.startswith("=HYPERLINK")
        finally:
            book.close()
    finally:
        session.close()


def test_cancel_export_publishes_nothing(convert, layout, tmp_path):
    session = convert(layout("veiculos"), ["descricao"], [["Sintético"]])
    target = tmp_path / "saida"
    with pytest.raises(Cancelled):
        export([session], target, Control(cancelled=lambda: True))
    assert not list(target.iterdir())


def test_orphan_items_block(convert, layout, tmp_path):
    head = convert(layout("os", "header"), ["Número da OS (externo)"], [["000001"]], name="cab.xlsx")
    item = convert(layout("os", "items"), ["Número da OS (externo)"], [["000002"]], name="item.xlsx")
    assert any(p.code == "ITENS_ORFAOS" for p in validate_links(head, item))
    with pytest.raises(ValueError, match="bloqueada"):
        export([head, item], tmp_path)


def test_all_twelve_layouts_export_synthetic(convert, layouts, tmp_path):
    sessions = []
    for index, layout in enumerate(item for item in layouts if item.path.parent.name == "V1"):
        values = {"UF": "SC"} if layout.family in {"clientes", "fornecedores"} else {"Tipo do item": "Serviço"} if layout.family == "produtos" else {"Número da OS (externo)": "000001"} if layout.family == "os" else {"Número da venda (externo)": "000001"} if layout.family == "vendas" else {layout.headers[0]: "Veículo sintético" if layout.family == "veiculos" else "Entidade sintética"}
        sessions.append(convert(layout, list(values), [list(values.values())], name=f"sintetico_{index}.xlsx"))
    for session in sessions:
        related = next((s for s in sessions if s.layout.family == session.layout.family and {session.layout.variant, s.layout.variant} == {"header", "items"}), None)
        out = export([session, related] if related else [session], tmp_path)
        assert list(out.glob("*.xlsx"))


def test_numeric_excel_decimal_not_ambiguous(convert, layout):
    session = convert(layout("produtos"), ["Tipo do item", "Preço à vista"], [["Serviço", 1.234]])
    assert session.stats()["errors"] == 0
    assert next(session.rows())["values"]["Preço à vista"] == "1,234"


def test_cli_restricted_arguments_and_json(monkeypatch, tmp_path):
    import io
    import app.ai.cli as cli
    monkeypatch.setattr(cli, "detect_clis", lambda c: [{"name": "claude", "enabled": True, "path": "/synthetic/claude"}])
    captured = []
    result = {"suggestions": []}
    class Fake:
        def __init__(self, args, **kwargs):
            captured.append((args, kwargs))
            self.stdin = io.StringIO()
            self.returncode = 0
        def communicate(self, timeout):
            return json.dumps({"structured_output": result}), ""
        def poll(self):
            return 0
    monkeypatch.setattr(cli.subprocess, "Popen", Fake)
    provider = CLIProvider(tmp_path)
    assert json.loads(provider.generate({"untrusted_data": "$(do not execute)"}, {})) == result
    args, kwargs = captured[0]
    assert kwargs["shell"] is False and args[args.index("--tools") + 1] == ""
    assert "--safe-mode" in args and "--strict-mcp-config" in args and "--no-session-persistence" in args
    assert "$(do not execute)" not in args
