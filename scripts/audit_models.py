"""Inventário estrutural sem reproduzir amostras pessoais."""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.core.catalog import digest, discover

root = Path(__file__).resolve().parents[1]
catalog = []
for layout in discover(root):
    book = load_workbook(layout.path)
    sheet = book[layout.sheet]
    columns = []
    for i, header in enumerate(layout.headers, 1):
        cells = [r[0] for r in sheet.iter_rows(min_row=2, min_col=i, max_col=i) if r[0].value not in (None, "")]
        columns.append({"position": i, "header": header, "types_observed": sorted({c.data_type for c in cells}), "formats_observed": sorted({c.number_format for c in cells}), "required": "somente quando declarado no Markdown; exemplos não comprovam obrigatoriedade"})
    catalog.append({"file": str(layout.path.relative_to(root)), "sha256": layout.digest, "sheet": layout.sheet, "family": layout.family, "variant": layout.variant, "dimensions": [sheet.max_row, sheet.max_column], "columns": columns, "validations": len(sheet.data_validations.dataValidation), "merges": list(map(str, sheet.merged_cells.ranges))})
    book.close()
(root / "docs/catalogo_modelos.json").write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
profile_path = root / "app/core/rules/profile.json"
if not profile_path.exists():
    profile_path.write_text(json.dumps({"version": 1, "source_sha256": digest(root / "INSTRUCOES_CONVERSOES.md"), "max_total_rows": 10000, "conflicts": ["COLUNAS_VAZIAS", "ATIVO_SIM_ATIVO", "TIPO_ITEM_DESCRICAO_CODIGO"], "source": "INSTRUCOES_CONVERSOES.md"}, indent=2) + "\n")
print(f"{len(catalog)} modelos auditados. Perfil existente nunca é atualizado automaticamente.")
