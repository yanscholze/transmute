"""Build nativo e pacote portátil, sem copiar estado, secrets ou exemplos.

Execute com o Python do .venv no SO alvo. PyInstaller não faz cross-compilation.
Cada release ganha diretório novo; versões anteriores não são apagadas.
"""
import argparse
import json
import os
import platform
import shutil
import subprocess
import sys
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.core.catalog import digest, discover, rules_path, verify_rules, workspace_config  # noqa: E402
from app import __version__  # noqa: E402


def archive_path(directory):
    """Append extension: version dots are not a filesystem suffix to replace."""
    return directory.with_name(directory.name + ".zip")


def validate_binary(command, release, env, timeout):
    """Surface frozen GUI tracebacks even when Windows has no console."""
    diagnostic = Path(env["CONVERSOR_STATE_DIR"]) / "build-error.txt"
    diagnostic.parent.mkdir(parents=True, exist_ok=True)
    check_env = {**env, "CONVERSOR_DIAGNOSTIC_FILE": str(diagnostic)}
    try:
        subprocess.run(command, cwd=release, env=check_env, check=True, timeout=timeout)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
        if diagnostic.is_file():
            print(diagnostic.read_text(encoding="utf-8"))
        raise


def build(library=None):
    system = platform.system()
    name = f"Transmute-{__version__}-{system}-{platform.machine()}-{datetime.now():%Y%m%d-%H%M%S}"
    release = ROOT / "dist" / name
    release.mkdir(parents=True, exist_ok=False)
    env = {**os.environ, "PYINSTALLER_CONFIG_DIR": str(ROOT / ".cache/pyinstaller"), "XDG_CACHE_HOME": str(ROOT / ".cache")}
    subprocess.run([sys.executable, "-m", "PyInstaller", "--noconfirm", "--clean", "--onefile", "--windowed", "--name", "ConversorDeBases", "--distpath", str(ROOT / "build/binary"), "--workpath", str(ROOT / "build/pyinstaller"), "--specpath", str(ROOT / "build"), "--add-data", str(ROOT / "app/core/rules/profile.json") + os.pathsep + "app/core/rules", "--collect-submodules", "keyring.backends", "--exclude-module", "PySide6.QtWebEngineCore", "--exclude-module", "PySide6.QtWebEngineWidgets", "--exclude-module", "PySide6.QtQml", "--exclude-module", "PySide6.QtQuick", str(ROOT / "run.py")], cwd=ROOT, env=env, check=True)
    executable = "ConversorDeBases.exe" if system == "Windows" else "ConversorDeBases"
    shutil.copy2(ROOT / "build/binary" / executable, release / executable)
    # Derive blank templates from the source structure; never distribute example rows.
    blank_cache, preset = {}, None
    if library:
        config = workspace_config(library)
        layouts = discover(library)
        profile = verify_rules(library)
        files = sorted({layout.path for layout in layouts})
        preset = {"brand": config.get("brand", "Conversor de Bases"), "primary_color": config.get("primary_color", "#09695D"), "theme_mode": config.get("theme_mode", "light"), "recognize_original_names": False, "families": {}}
        logo = Path(config.get("logo_path", ""))
        if logo.is_file():
            relative_logo = "marca/logo" + logo.suffix.lower()
            target_logo = release / relative_logo
            target_logo.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(logo, target_logo)
            preset.update(logo_path=relative_logo, logo_size=config.get("logo_size", 50),
                          logo_opacity=config.get("logo_opacity", 8), logo_position=config.get("logo_position", "center"))
    else:
        files = sorted((ROOT / "#PLANILHAS DE BASE").rglob("*.xlsx"))
    parents = sorted({path.parent for path in files})
    counts = Counter(parent.name for parent in parents)
    group_names = {parent: parent.name if counts[parent.name] == 1 else f"Conjunto_{index+1:03d}_{parent.name}" for index, parent in enumerate(parents)}
    for source in files:
        relative = Path("#PLANILHAS DE BASE") / group_names[source.parent] / source.name if library else source.relative_to(ROOT)
        target = release / relative
        if preset is not None:
            preset["families"][relative.as_posix()] = next(layout.family for layout in layouts if layout.path == source)
        target.parent.mkdir(parents=True, exist_ok=True)
        signature = digest(source)
        if signature in blank_cache:
            shutil.copy2(blank_cache[signature], target)
            continue
        workbook = load_workbook(source)
        try:
            for sheet in workbook:
                if sheet.max_row > 1:
                    sheet.delete_rows(2, sheet.max_row - 1)
            workbook.save(target)
        finally:
            workbook.close()
        blank_cache[signature] = target
    shutil.copy2(rules_path(library) if library else ROOT / "INSTRUCOES_CONVERSOES.md", release / "INSTRUCOES_CONVERSOES.md")
    if preset is not None:
        (release / "regras_perfil.json").write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
        (release / "library_preset.json").write_text(json.dumps(preset, ensure_ascii=False, indent=2), encoding="utf-8")
    else:
        shutil.copy2(ROOT / "#PLANILHAS DE BASE/VERSOES.md", release / "#PLANILHAS DE BASE/VERSOES.md")
        shutil.copy2(ROOT / "packaging/brand.json", release / "library_preset.json")
    shutil.copytree(ROOT / "docs", release / "docs", ignore=shutil.ignore_patterns("*.png", "originais.sha256.json", "catalogo_modelos.json"))
    (release / "LEIA-ME.txt").write_text(f"Transmute {__version__}\n\nExtraia todo o ZIP em pasta gravável e abra " + executable + ".\nPython e bibliotecas estão incluídos. Mantenha regras e modelos ao lado do app.\nDashboard reúne atalhos, trabalhos e indicadores. O Estúdio de Conversão reúne fontes, regras, conversa com IA, planos aprovados e auditoria.\nBiblioteca contém modelos/regras; Aparência contém temas claro/escuro, destaque e logo.\nOs modelos deste pacote são cópias vazias, sem exemplos. Fontes originais do projeto não foram alteradas.\nEstado portátil fica em .local; não inclua essa pasta ao redistribuir.\nConsulte docs para regras, privacidade, manutenção e validação.\n", encoding="utf-8")
    licenses = release / "licencas"
    licenses.mkdir()
    import importlib.metadata
    for dist in importlib.metadata.distributions():
        for entry in dist.files or []:
            if any(s in str(entry).lower() for s in ("license", "copying", "notice")):
                source = Path(dist.locate_file(entry))
                if source.is_file():
                    dest = licenses / (dist.metadata["Name"] + "_" + str(entry).replace("/", "_").replace("\\", "_"))
                    shutil.copy2(source, dest)
    # Test the packed executable with no project interpreter dependency.
    test_env = {**env, "CONVERSOR_STATE_DIR": str(ROOT / ".local/build-validation" / name), "QT_QPA_PLATFORM": "offscreen"}
    if not library:
        validate_binary([str(release / executable), "--self-test", "--self-test-report", str(Path(test_env["CONVERSOR_STATE_DIR"]) / "selftest.json")], release, test_env, 120)
    validate_binary([str(release / executable), "--smoke-test"], release, test_env, 60)
    # The first onefile process deletes _MEI on exit. Reopening the same state
    # proves no active rule profile still points into that temporary directory.
    validate_binary([str(release / executable), "--smoke-test"], release, test_env, 60)
    manifest = {str(p.relative_to(release)): digest(p) for p in release.rglob("*") if p.is_file()}
    (release / "SHA256.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    archive = archive_path(release)
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zip:
        for path in sorted(release.rglob("*")):
            if path.is_file():
                zip.write(path, Path(release.name) / path.relative_to(release))
    with zipfile.ZipFile(archive) as zip:
        if zip.testzip() is not None:
            raise RuntimeError("ZIP corrompido.")
    print(f"PACOTE VALIDADO: {archive}")
    return archive


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build nativo, opcionalmente com biblioteca white label")
    parser.add_argument("--library", type=Path, help="Diretório contendo workspace.json da biblioteca aprovada")
    build(parser.parse_args().library)
