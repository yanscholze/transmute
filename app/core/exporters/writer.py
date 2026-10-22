"""Exportação textual e transacional: planejar antes, publicar só ao terminar.

Partes vinculadas compartilham grupos; quando preservar ordem e vínculo é
impossível, bloquear é mais seguro do que reorganizar silenciosamente.
"""
import json
import shutil
import tempfile
import uuid
from collections import Counter
from copy import copy
from dataclasses import asdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from app.core.catalog import verify_rules
from app.core.mapping import targets
from app.core.session import validate_links
from app.models import Control


def selected_lines(session, partial):
    if not partial:
        return None
    # Distinct scenarios first, then fill remaining slots with distinct records/groups.
    candidates, scenarios, chosen_keys, remaining = [], set(), set(), []
    for row in session.rows():
        v = row["values"]
        key = row["group"] or str(row["line"])
        if key in chosen_keys:
            continue
        doc = v.get("CPF/CNPJ", "")
        scenario = ("CPF" if len(doc) == 11 else "CNPJ" if len(doc) == 14 else "SEM_DOC", bool(v.get("Inscrição Estadual")), row["severity"], v.get("Tipo do item", ""))
        if scenario not in scenarios and len(candidates) < 5:
            candidates.append(row["line"])
            chosen_keys.add(key)
            scenarios.add(scenario)
        else:
            remaining.append((row["line"], key))
    for line, key in remaining:
        if len(candidates) >= 5:
            break
        if key not in chosen_keys:
            candidates.append(line)
            chosen_keys.add(key)
    if session.layout.family in {"os", "vendas"}:
        return {row["line"] for row in session.rows() if row["group"] in chosen_keys}
    return set(candidates)


def plan_parts(sessions, selections, max_records):
    """Planeje partes por grupo e prove monotonicidade em cada fonte."""
    primary = sessions[0]
    grouped = primary.layout.family in {"os", "vendas"}
    counts = [Counter() for _ in sessions]
    order = []
    seen = set()
    for index, session in enumerate(sessions):
        for row in session.rows():
            if selections[index] is not None and row["line"] not in selections[index]:
                continue
            key = row["group"] if grouped else str(row["line"])
            counts[index][key] += 1
            if key not in seen:
                seen.add(key)
                order.append(key)
    assignment, part, totals = {}, 1, [0] * len(sessions)
    for key in order:
        if any(c[key] > max_records for c in counts):
            raise ValueError("Um único grupo excede 9.999 registros. Não é possível dividi-lo sem violar os vínculos.")
        if any(totals[i] + c[key] > max_records for i, c in enumerate(counts)):
            part += 1
            totals = [0] * len(sessions)
        assignment[key] = part
        totals = [totals[i] + c[key] for i, c in enumerate(counts)]
    # Guarantee that partitioning does not silently reorder records.
    for i, session in enumerate(sessions):
        previous = 0
        for row in session.rows():
            if selections[i] is not None and row["line"] not in selections[i]:
                continue
            key = row["group"] if grouped else str(row["line"])
            current = assignment[key]
            if current < previous:
                raise ValueError("Grupos intercalados/ordens divergentes impedem divisão preservando a ordem. Revise a ordenação na origem e reprocesse; nada foi reordenado.")
            previous = current
    return assignment, part


def write_workbook(path, session, rows, columns, control):
    """Derive cabeçalhos/estilos do modelo e escreva dados como texto literal."""
    book = Workbook(write_only=True)
    sheet = book.create_sheet(session.layout.sheet)
    model = load_workbook(session.layout.path, read_only=False)
    try:
        template = model[session.layout.sheet]
        model_headers = session.layout.headers
        head = []
        for field in columns:
            c = WriteOnlyCell(sheet, field)
            c.data_type = "s"
            if field in model_headers:
                c._style = copy(template.cell(1, model_headers.index(field) + 1)._style)
            head.append(c)
        sheet.append(head)
        for row in rows:
            control.check()
            cells = []
            for field in columns:
                value = row["values"].get(field, "")
                cell = WriteOnlyCell(sheet, value if value != "" else None)
                if value != "":
                    cell.data_type = "s"  # includes leading '=', '+', '@': never execute source content
                    cell.number_format = "@"
                cells.append(cell)
            sheet.append(cells)
        book.save(path)
    finally:
        model.close()
        book.close()


def report_data(sessions, selections, issues, outputs=None, omitted=None):
    datasets = []
    for i, session in enumerate(sessions):
        unresolved = [h for h in targets(session.layout) if not any(d.target == h for d in session.decisions)]
        counts = Counter()
        problematic = []
        for row in session.rows():
            for change in row["changes"]:
                counts[change["rule"]] += 1
            if row["issues"]:
                problematic.append({"line": row["line"], "issues": row["issues"]})
        options = asdict(session.options)
        options.pop("overrides", None)
        datasets.append({"original": session.source.path.name, "source_sha256": session.source_hash,
                         "sheet": session.source.sheet, "header_row": session.source.header_row,
                         "encoding": session.source.encoding, "delimiter": session.source.delimiter,
                         "layout": session.layout.label, "template": session.layout.path.name,
                         "template_sha256": session.layout.digest, "rules_sha256": session.rules_hash,
                         "profile_sha256": session.profile_hash,
                         "statistics": session.stats(),
                         "selected": len(selections[i]) if selections[i] is not None else session.stats()["read"],
                         "partial_selection": "Cenários distintos priorizados, até cinco registros/grupos. Cobertura não garante todos os cenários existentes." if selections[i] is not None else "Base completa",
                         "mapped": [asdict(d) for d in session.decisions if d.target],
                         "source_headers": dict(zip(session.source.ids, session.source.headers)),
                         "ignored": [h for sid, h in zip(session.source.ids, session.source.headers) if not any(d.source == sid and d.target for d in session.decisions)],
                         "unresolved_fields": unresolved, "transformations": dict(counts),
                         "ai_decisions": session.ai_events, "problems": problematic,
                         "global_issues": [asdict(p) for p in session.global_issues], "options": options})
    return {"application": "Conversor de Bases 1.0.0", "created": datetime.now().astimezone().isoformat(),
            "datasets": datasets, "link_issues": [asdict(p) for p in issues], "outputs": outputs or [], "omitted_columns": omitted or {},
            "scope": "Validação das regras documentadas; requisitos não documentados e aceitação em ERP exigem teste de importação."}


def save_reports(directory, report):
    (directory / "relatorio.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["# Relatório da conversão", "", f"Data: {report['created']}", "", report["scope"], ""]
    for data in report["datasets"]:
        stats = data["statistics"]
        lines += [f"## {data['layout']}", f"Origem: {data['original']} / aba {data['sheet']}", f"Lidos: {stats['read']}; válidos sem avisos: {stats['valid']}; com avisos: {stats['warnings']}; com erros: {stats['errors']}; linhas vazias: {stats['blank_rows']}; selecionados: {data['selected']}.", data["partial_selection"],
                  "Colunas ignoradas: " + ", ".join(data["ignored"]), "Campos sem origem mapeada (podem ser opcionais ou preenchidos por regra): " + ", ".join(data["unresolved_fields"]), "Transformações: " + json.dumps(data["transformations"], ensure_ascii=False), ""]
        for p in data["global_issues"]:
            lines.append(f"- {p['severity']} {p['code']}: {p['message']}")
        for row in data["problems"]:
            for p in row["issues"]:
                lines.append(f"- Linha {row['line']} — {p['severity']} — {p['field']} — {p['code']}: {p['message']}")
    lines += ["", "## Vínculos"] + [f"- {p['severity']}: {p['message']}" for p in report["link_issues"]]
    lines += ["", "## Colunas omitidas", json.dumps(report["omitted_columns"], ensure_ascii=False), "", "## Arquivos", *report["outputs"]]
    (directory / "diagnostico.md").write_text("\n".join(lines), encoding="utf-8")
    summary = [f"Conversor de Bases — {report['created']}"]
    for data in report["datasets"]:
        summary.append(f"{data['layout']}: {json.dumps(data['statistics'], ensure_ascii=False)}; selecionados {data['selected']}")
    summary += ["Colunas omitidas: " + json.dumps(report["omitted_columns"], ensure_ascii=False), "Consulte diagnostico.md e relatorio.json para problemas, fontes e decisões."]
    (directory / "resumo.txt").write_text("\n".join(summary), encoding="utf-8")


def export(sessions, destination, control=None, reports_only=False, detailed_audit=False, diagnostics_destination=None):
    """Nunca sobrescreva fontes/saídas; em erro ou cancelamento retire o staging."""
    control = control or Control()
    for session in sessions:
        session.verify()
    linked_issues = validate_links(sessions[0], sessions[1] if len(sessions) > 1 else None)
    if not reports_only and (any(s.stats()["errors"] or any(p.severity == "ERRO" for p in s.global_issues) for s in sessions) or any(p.severity == "ERRO" for p in linked_issues)):
        raise ValueError("Exportação final bloqueada por erros. Corrija e revalide, ou exporte somente o relatório.")
    selections = [selected_lines(s, s.options.partial) for s in sessions]
    if len(sessions) == 2 and sessions[0].options.partial:
        groups = {r["group"] for r in sessions[0].rows() if r["line"] in selections[0]}
        selections[1] = {r["line"] for r in sessions[1].rows() if r["group"] in groups}
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    final = destination / ("Conversao_" + datetime.now().strftime("%Y%m%d_%H%M%S_") + uuid.uuid4().hex[:8])
    staging = Path(tempfile.mkdtemp(prefix=".conversor-", dir=destination))
    outputs, omitted = [], {}
    diagnostic_staging, diagnostic_final = None, None
    published = False
    try:
        if not reports_only:
            profile = verify_rules(sessions[0].root)
            assignment, parts = plan_parts(sessions, selections, profile["max_total_rows"] - 1)
            for i, session in enumerate(sessions):
                present = set()
                for row in session.rows():
                    if selections[i] is None or row["line"] in selections[i]:
                        present.update(k for k, value in row["values"].items() if value != "")
                columns = [h for h in targets(session.layout) if session.options.columns == "full" or h in present]
                omitted[session.layout.label] = [h for h in targets(session.layout) if h not in columns]
                for part in range(1, parts + 1):
                    control.check()
                    name = f"{session.layout.family}_{session.layout.variant}_parte_{part:03d}.xlsx"
                    def rows_for_part(s=session, idx=i, p=part):
                        for row in s.rows():
                            if selections[idx] is not None and row["line"] not in selections[idx]:
                                continue
                            key = row["group"] if s.layout.family in {"os", "vendas"} else str(row["line"])
                            if assignment[key] == p:
                                yield row
                    write_workbook(staging / name, session, rows_for_part(), columns, control)
                    outputs.append(name)
                    control.progress(part, f"Gravando parte {part}/{parts}")
        report = report_data(sessions, selections, linked_issues, outputs, omitted)
        save_reports(staging, report)
        if detailed_audit:
            with (staging / "auditoria_sensivel.jsonl").open("w", encoding="utf-8") as stream:
                for session in sessions:
                    for row in session.rows():
                        control.check()
                        stream.write(json.dumps({"layout": session.layout.label, **row}, ensure_ascii=False) + "\n")
        control.check()
        for session in sessions:
            session.verify()
        if diagnostics_destination is not None and Path(diagnostics_destination).resolve() != destination.resolve():
            diag = Path(diagnostics_destination)
            diag.mkdir(parents=True, exist_ok=True)
            diagnostic_staging = Path(tempfile.mkdtemp(prefix=".conversor-", dir=diag))
            diagnostic_final = diag / final.name
            for name in ["diagnostico.md", "relatorio.json", "resumo.txt"]:
                shutil.copy2(staging / name, diagnostic_staging / name)
        staging.rename(final)
        published = True
        if diagnostic_staging:
            diagnostic_staging.rename(diagnostic_final)
        return final
    except BaseException:
        shutil.rmtree(staging, ignore_errors=True)
        if diagnostic_staging:
            shutil.rmtree(diagnostic_staging, ignore_errors=True)
        if published:
            shutil.rmtree(final, ignore_errors=True)
        raise
