"""Estrutura vem dos modelos; comportamento vem de um perfil revisado.

Uma biblioteca pode apontar a fontes externas ou a cópias importadas. Exemplos
de dados não são retornados pelo catálogo e nunca viram valores padrão.
"""
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from app.models import Layout


def digest(path):
    """Hash incremental usado para provar que as fontes não mudaram."""
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def discover(root):
    """Leia todas as abas/modelos e preserve cabeçalhos literais e sua ordem."""
    layouts = []
    config = workspace_config(root)
    base = Path(config.get("models_dir", Path(root) / "#PLANILHAS DE BASE"))
    files = set(base.rglob("*.xlsx")) if base.is_dir() else set()
    files.update(Path(p) for p in config.get("uploaded_models", []) if Path(p).is_file())
    for path in sorted(files):
        name = path.name
        family = ("clientes" if name.startswith("clientes") else "fornecedores" if name.startswith("fornecedores")
                  else "produtos" if name.startswith("produtos") else "veiculos" if name.startswith("cad_veiculo")
                  else "pagar" if "pagar" in name else "receber" if "receber" in name
                  else "os" if name.startswith("ordens") else "vendas" if name.startswith("vendas") else "")
        family = config.get("families", {}).get(str(path), family if config.get("recognize_original_names", True) else "generic") or "generic"
        variant = "items" if "somente_itens" in name else "header" if "somente_cabe" in name else "joined" if "juntos" in name else "record"
        if family == "generic":
            variant = "record"
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in book:
                row = next(sheet.iter_rows(max_row=1, values_only=True))
                headers = [str(v) if v is not None else "" for v in row]
                if not all(headers) or len(headers) != len(set(headers)):
                    raise ValueError(f"Modelo com cabeçalho vazio ou duplicado: {path.name}")
                layouts.append(Layout(path, sheet.title, headers, family, variant, digest(path)))
        finally:
            book.close()
    return layouts


def verify_rules(root):
    """Recuse prosa alterada sem uma revisão correspondente do perfil."""
    path = rules_path(root)
    config = workspace_config(root)
    profile = json.loads(Path(config.get("profile_path", Path(__file__).parent / "rules" / "profile.json")).read_text(encoding="utf-8"))
    if not path.is_file() or digest(path) != profile["source_sha256"]:
        raise ValueError("O arquivo de regras está ausente ou mudou desde a revisão. Abra Regras na tela inicial, revise o texto e as regras executáveis e salve uma revisão aprovada.")
    from app.core.rules.declarative import validate_profile
    validate_profile(profile)
    return profile


def workspace_config(root):
    path = Path(root) / "workspace.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}


def rules_path(root):
    return Path(workspace_config(root).get("rules_path", Path(root) / "INSTRUCOES_CONVERSOES.md"))
