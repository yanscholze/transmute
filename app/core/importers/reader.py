"""Detecção por conteúdo e leitores que expõem posição física e tipo de origem.

CSV e XLSX são iterados; HTML tem limite explícito de tamanho. Nenhum leitor
executa fórmula ou tenta recuperar zeros que já desapareceram da fonte.
"""
import csv
import re
from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal
from html.parser import HTMLParser
from itertools import islice
from pathlib import Path

from charset_normalizer import from_bytes
from openpyxl import load_workbook

from app.models import Cell, Control, Source


def excel_cell(value, kind="", fmt="General"):
    """Converta representação sem float arithmetic, mantendo evidência de risco."""
    if value is None:
        return Cell()
    if kind == "f":
        return Cell(str(value), "formula", "Fórmula na origem; forneça valor revisado, sem executar a fórmula.")
    if kind == "e":
        return Cell(str(value), "error", "Erro de célula na origem.")
    if isinstance(value, bool):
        return Cell("TRUE" if value else "FALSE", "boolean")
    if isinstance(value, datetime):
        return Cell(value.strftime("%d/%m/%Y %H:%M:%S") if value.time() != time() else value.strftime("%d/%m/%Y"), "date")
    if isinstance(value, date):
        return Cell(value.strftime("%d/%m/%Y"), "date")
    if isinstance(value, time):
        return Cell(value.strftime("%H:%M:%S"), "time")
    if isinstance(value, (int, float)):
        text = format(Decimal(str(value)), "f")
        if "." in text:
            text = text.rstrip("0").rstrip(".")
        if re.fullmatch(r"0+", fmt) and float(value).is_integer():
            return Cell(text.zfill(len(fmt)), "formatted_number", "Zeros preservados pelo formato explícito do Excel.")
        return Cell(text, "number", "Valor numérico na origem; zeros anteriores já podem ter sido perdidos.")
    return Cell(str(value))


def detect(path):
    path = Path(path)
    with path.open("rb") as stream:
        head = stream.read(8192)
    if head.startswith(b"PK\x03\x04"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if re.search(br"<(?:html|table)\b", head.lstrip(b"\xef\xbb\xbf \r\n\t"), re.I):
        return "html"
    if path.suffix.lower() not in {".csv", ".tsv", ".txt", ".xls"}:
        raise ValueError("Formato não reconhecido. Use XLSX, XLS ou CSV. PDF ainda não é suportado.")
    return "csv"


def csv_settings(path):
    with Path(path).open("rb") as stream:
        sample = stream.read(128 * 1024)
    if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    else:
        try:
            sample.decode("utf-8-sig")
            encoding = "utf-8-sig"
        except UnicodeDecodeError:
            match = from_bytes(sample).best()
            # Legacy Brazilian exports commonly use Windows-1252. Short samples
            # cannot distinguish it reliably from other Latin code pages.
            # Prefer that explicit locale default and expose an override in UI.
            try:
                sample.decode("cp1252")
                encoding = "cp1252"
            except UnicodeDecodeError:
                encoding = match.encoding if match else "cp1252"
    decoded = sample.decode(encoding, errors="strict")
    try:
        delimiter = csv.Sniffer().sniff(decoded, delimiters=";,\t|").delimiter
    except csv.Error:
        delimiter = max(";,\t|", key=lambda d: sum(line.count(d) for line in decoded.splitlines()[:20]))
    return encoding, delimiter


class HTMLTables(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables, self.table, self.row, self.value = [], None, None, None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            if self.table is not None:
                raise ValueError("HTML com tabelas aninhadas requer revisão externa.")
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            if any(k in {"colspan", "rowspan"} and v != "1" for k, v in attrs):
                raise ValueError("HTML com células mescladas requer revisão externa.")
            self.value = ""
        elif tag == "br" and self.value is not None:
            self.value += "\n"

    def handle_data(self, data):
        if self.value is not None:
            self.value += data

    def handle_endtag(self, tag):
        if tag in {"td", "th"} and self.value is not None:
            self.row.append(Cell(self.value))
            self.value = None
        elif tag == "tr" and self.row is not None:
            self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            self.tables.append(self.table)
            self.table = None


def html_tables(path):
    if Path(path).stat().st_size > 64 * 1024 * 1024:
        raise ValueError("XLS HTML acima de 64 MB: salve como XLSX/CSV para leitura incremental.")
    encoding, _ = csv_settings(path)
    parser = HTMLTables()
    parser.feed(Path(path).read_text(encoding=encoding))
    return parser.tables


def sheets(path):
    fmt = detect(path)
    if fmt == "xlsx":
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            return book.sheetnames
        finally:
            book.close()
    if fmt == "xls":
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            return book.sheet_names()
        finally:
            book.release_resources()
    if fmt == "html":
        return [f"Tabela {i + 1}" for i in range(len(html_tables(path)))]
    return ["CSV"]


def raw_rows(path, sheet, fmt, encoding="", delimiter="", control=None):
    control = control or Control()
    if fmt == "xlsx":
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            for n, row in enumerate(book[sheet].iter_rows(), 1):
                control.check()
                yield n, [excel_cell(c.value, c.data_type, c.number_format) for c in row]
        finally:
            book.close()
    elif fmt == "xls":
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True, formatting_info=True)
        try:
            tab = book.sheet_by_name(sheet)
            for n in range(tab.nrows):
                control.check()
                row = []
                for c in tab.row(n):
                    value = xlrd.xldate_as_datetime(c.value, book.datemode) if c.ctype == xlrd.XL_CELL_DATE else c.value
                    fmt = book.format_map[book.xf_list[c.xf_index].format_key].format_str
                    row.append(excel_cell(value, "e" if c.ctype == xlrd.XL_CELL_ERROR else "", fmt))
                yield n + 1, row
        finally:
            book.release_resources()
    elif fmt == "html":
        for n, row in enumerate(html_tables(path)[int(sheet.split()[-1]) - 1], 1):
            control.check()
            yield n, row
    else:
        csv.field_size_limit(4 * 1024 * 1024)
        with Path(path).open(encoding=encoding, newline="", errors="strict") as stream:
            reader = csv.reader(stream, delimiter=delimiter, strict=True)
            for row in reader:
                control.check()
                yield reader.line_num, [Cell(v) for v in row]


def inspect(path, sheet=None, header_row=None, encoding=None, delimiter=None, control=None):
    """Sugira estrutura a partir de amostra limitada; a UI permite corrigir a escolha."""
    path = Path(path).resolve()
    fmt = detect(path)
    available = sheets(path)
    if not available:
        raise ValueError("Nenhuma aba/tabela encontrada.")
    sheet = sheet or available[0]
    enc, delim = csv_settings(path) if fmt == "csv" else ("", "")
    enc, delim = encoding or enc, delimiter or delim
    candidates = list(islice(raw_rows(path, sheet, fmt, enc, delim, control), max(50, header_row or 0)))
    nonempty = [(n, row) for n, row in candidates if any(c.text != "" for c in row)]
    if not nonempty:
        raise ValueError("Arquivo/aba sem dados.")
    if header_row is None:
        def score(item):
            _, row = item
            values = [c.text.strip() for c in row if c.text.strip()]
            words = sum(bool(re.search(r"[A-Za-zÀ-ÿ]", v)) for v in values)
            return len(values) * 2 + words + len(set(values)) * .2
        header_row, row = max(nonempty, key=score)
    else:
        row = next((r for n, r in candidates if n == header_row), None)
        if row is None:
            raise ValueError("Linha de cabeçalho não encontrada.")
    while row and row[-1].text == "":
        row = row[:-1]
    headers = [c.text for c in row]
    warnings = [f"Cabeçalho duplicado: {h}" for h, count in Counter(headers).items() if count > 1]
    if any(not h.strip() for h in headers):
        warnings.append("Há colunas sem cabeçalho. Use os IDs de coluna para mapear.")
    if fmt == "csv":
        warnings.append(f"Encoding sugerido: {enc}; confirme acentos e delimitador na prévia.")
    return Source(path, sheet, header_row, headers, enc, delim, fmt, warnings)


def records(source, control=None):
    """Inclua linhas vazias e células excedentes para contagem e auditoria explícitas."""
    for n, row in raw_rows(source.path, source.sheet, source.format, source.encoding, source.delimiter, control):
        if n <= source.header_row:
            continue
        if not any(c.text != "" for c in row):
            yield n, None, []
            continue
        extra = row[len(source.headers):]
        problems = []
        if any(c.text != "" for c in extra) or (source.format in {"csv", "html"} and len(row) != len(source.headers)):
            problems.append("Quantidade de células difere do cabeçalho; nenhuma célula extra foi descartada da auditoria.")
        padded = row + [Cell()] * max(0, len(source.headers) - len(row))
        yield n, {f"C{i + 1}": cell for i, cell in enumerate(padded)}, problems
