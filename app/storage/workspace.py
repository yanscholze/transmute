"""Biblioteca configurável. Referências externas são somente leitura.

Uploads e edições são copiados para o estado local. Cada revisão de regras é
um novo par Markdown/JSON; o ponteiro só muda depois que ambos são gravados.
"""
import json
import re
import shutil
import uuid
from pathlib import Path

from app.core.catalog import digest, workspace_config


def save_config(root, config):
    path = Path(root) / "workspace.json"
    temp = path.with_suffix(".new")
    temp.write_text(json.dumps(config, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temp.replace(path)


def initialize(project, state):
    root = state / "workspace"
    root.mkdir(parents=True, exist_ok=True, mode=0o700)
    if not (root / "workspace.json").exists():
        preset_path = project / "library_preset.json"
        preset = json.loads(preset_path.read_text(encoding="utf-8")) if preset_path.is_file() else {}
        profile = Path(__file__).resolve().parents[1] / "core/rules/profile.json"
        # A portable directory may contain its own approved declarative profile.
        if (project / "regras_perfil.json").is_file():
            profile = project / "regras_perfil.json"
        families = {str(project / key): value for key, value in preset.get("families", {}).items()}
        save_config(root, {"models_dir": str(project / "#PLANILHAS DE BASE"), "uploaded_models": [], "rules_path": str(project / "INSTRUCOES_CONVERSOES.md"), "profile_path": str(profile), "recognize_original_names": preset.get("recognize_original_names", True), "families": families, "brand": preset.get("brand", "Conversor de Bases"), "primary_color": preset.get("primary_color", "#09695D")})
        config = workspace_config(root)
        config.update(logo_path=str(project / preset["logo_path"]) if preset.get("logo_path") else "",
                      logo_size=preset.get("logo_size", 50), logo_opacity=preset.get("logo_opacity", 8), logo_position=preset.get("logo_position", "center"))
        save_config(root, config)
    return root


def import_model(root, source):
    source = Path(source)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("Modelos de destino precisam ser XLSX.")
    # Validate before committing a library pointer; a bad upload must be recoverable.
    from openpyxl import load_workbook
    workbook = load_workbook(source, read_only=True)
    try:
        for sheet in workbook:
            headers = next(sheet.iter_rows(max_row=1, values_only=True))
            if any(v is None or str(v) == "" for v in headers) or len(headers) != len(set(headers)):
                raise ValueError("Modelo possui cabeçalho vazio ou duplicado; ajuste uma cópia antes de importar.")
    finally:
        workbook.close()
    folder = Path(root) / "models" / uuid.uuid4().hex[:10]
    folder.mkdir(parents=True, exist_ok=False)
    destination = folder / source.name
    shutil.copy2(source, destination)
    config = workspace_config(root)
    config.setdefault("uploaded_models", []).append(str(destination))
    config.setdefault("families", {})[str(destination)] = "generic"
    save_config(root, config)
    return destination


def save_rule_revision(root, text, profile):
    from app.core.rules.declarative import validate_profile
    validate_profile(profile)
    if len(text) > 2 * 1024 * 1024:
        raise ValueError("Arquivo de regras acima de 2 MB.")
    folder = Path(root) / "rule-revisions" / uuid.uuid4().hex
    folder.mkdir(parents=True, exist_ok=False, mode=0o700)
    rules = folder / "regras.md"
    rules.write_text(text, encoding="utf-8")
    profile = {**profile, "source_sha256": digest(rules)}
    target = folder / "perfil.json"
    target.write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
    config = workspace_config(root)
    config.update(rules_path=str(rules), profile_path=str(target))
    save_config(root, config)
    return rules


def set_brand(root, name, color):
    if not name.strip() or len(name) > 60 or not re.fullmatch(r"#[0-9A-Fa-f]{6}", color):
        raise ValueError("Nome precisa ter 1–60 caracteres e cor deve ser hexadecimal #RRGGBB.")
    config = workspace_config(root)
    config.update(brand=name.strip(), primary_color=color)
    save_config(root, config)
